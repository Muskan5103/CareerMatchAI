
import re
import zipfile

from docx import Document
from pypdf import PdfReader


MAX_EXTRACTED_TEXT = 200_000

SKILL_TERMS = [
    "python", "django", "flask", "fastapi", "java", "javascript",
    "typescript", "react", "node.js", "html", "css", "bootstrap",
    "sql", "postgresql", "mysql", "mongodb", "sqlite",
    "machine learning", "deep learning", "natural language processing",
    "nlp", "scikit-learn", "pandas", "numpy", "tensorflow", "pytorch",
    "data analysis", "data visualization", "power bi", "tableau",
    "excel", "statistics", "computer vision", "opencv", "yolo",
    "git", "github", "docker", "kubernetes", "aws", "azure",
    "rest api", "api development", "unit testing", "pytest",
    "selenium", "software testing", "agile", "communication",
    "project management", "problem solving", "c++", "c#", ".net",
]


def extract_pdf_text(file_obj):
    file_obj.seek(0)
    reader = PdfReader(file_obj, strict=False)

    if reader.is_encrypted:
        raise ValueError(
            "Password-protected PDFs are not supported."
        )

    chunks = []

    for page in reader.pages:
        text = page.extract_text() or ""
        chunks.append(text)

        if sum(len(chunk) for chunk in chunks) > MAX_EXTRACTED_TEXT:
            break

    return "\n".join(chunks)[:MAX_EXTRACTED_TEXT]




def extract_docx_text(file_obj):
    file_obj.seek(0)
    document = Document(file_obj)

    chunks = [paragraph.text for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            chunks.append(" ".join(cell.text for cell in row.cells))

    return "\n".join(chunks)[:MAX_EXTRACTED_TEXT]


def extract_resume_text(file_obj, filename):
    filename = filename.lower()

    try:
        if filename.endswith(".pdf"):
            text = extract_pdf_text(file_obj)
        elif filename.endswith(".docx"):
            text = extract_docx_text(file_obj)
        else:
            raise ValueError("Unsupported resume format.")
    except Exception as exc:
        if isinstance(exc, ValueError):
            raise
        raise ValueError(
            "Could not read this document. Please upload a valid PDF or DOCX."
        ) from exc

    if not text.strip():
        raise ValueError(
            "No readable text was found. Scanned PDFs may need OCR."
        )

    return text



import re


def extract_skills(text):
    text_lower = text.lower()
    found = []

    for skill in SKILL_TERMS:
        pattern = (
            r"(?<!\w)"
            + re.escape(skill.lower())
            + r"(?!\w)"
        )

        if re.search(pattern, text_lower):
            found.append(skill)

    return sorted(set(found))




def normalize_skill_list(skills):
    if not skills:
        return []

    if isinstance(skills, str):
        skills = skills.split(",")

    return sorted({
        str(skill).strip().lower()
        for skill in skills
        if str(skill).strip()
    })




from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


def calculate_job_match(resume, job):
    resume_skills = normalize_skill_list(resume.extracted_skills)
    required_skills = normalize_skill_list(job.required_skills)

    matched_skills = sorted(set(resume_skills) & set(required_skills))
    missing_skills = sorted(set(required_skills) - set(resume_skills))

    if required_skills:
        skill_coverage = len(matched_skills) / len(required_skills)
    else:
        skill_coverage = 0.0

    resume_text = " ".join([
        resume.extracted_text or "",
        " ".join(resume_skills),
    ])

    job_text = " ".join([
        job.title or "",
        job.description or "",
        job.company_name or "",
        " ".join(required_skills),
    ])

    try:
        vectorizer = TfidfVectorizer(
            lowercase=True,
            stop_words="english",
            ngram_range=(1, 2),
            max_features=10000,
        )

        vectors = vectorizer.fit_transform([resume_text, job_text])
        text_similarity = float(
            cosine_similarity(vectors[0:1], vectors[1:2])[0][0]
        )
    except ValueError:
        text_similarity = 0.0

    raw_score = (
        70 * text_similarity
        + 30 * skill_coverage
    )

    return {
        "score": round(max(0.0, min(100.0, raw_score)), 2),
        "text_similarity": round(text_similarity * 100, 2),
        "skill_coverage": round(skill_coverage * 100, 2),
        "matched_skills": matched_skills,
        "missing_skills": missing_skills,
    }


def recommend_jobs(resume, jobs, limit=20):
    recommendations = []

    for job in jobs:
        match = calculate_job_match(resume, job)
        recommendations.append({
            "job": job,
            **match,
        })

    recommendations.sort(
        key=lambda item: item["score"],
        reverse=True,
    )

    return recommendations[:limit]
